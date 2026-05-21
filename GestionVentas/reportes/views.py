from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.contrib import messages

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from clientes.models import Cliente
from ventas.models import OportunidadVenta, Seguimiento

from users.decorators import admin_required


@login_required
@admin_required
def reporte_clientes_pdf(request):

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_clientes.pdf"'

    pdf = canvas.Canvas(response, pagesize=letter)
    width, height = letter

    pdf.setTitle("Reporte de Clientes")

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, height - 50, "Reporte de Clientes")

    pdf.setFont("Helvetica", 10)
    pdf.drawString(50, height - 70, "Sistema GestionVentas CRM")

    y = height - 110

    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(50, y, "ID")
    pdf.drawString(80, y, "Nombre")
    pdf.drawString(210, y, "Documento")
    pdf.drawString(310, y, "Correo")
    pdf.drawString(470, y, "Estado")

    y -= 20
    pdf.setFont("Helvetica", 8)

    clientes = Cliente.objects.all().order_by('nombre')

    for cliente in clientes:

        if y < 50:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 8)

        pdf.drawString(50, y, str(cliente.id))
        pdf.drawString(80, y, str(cliente.nombre)[:22])
        pdf.drawString(210, y, str(cliente.documento)[:18])
        pdf.drawString(310, y, str(cliente.correo)[:28])
        pdf.drawString(470, y, str(cliente.estado).title())

        y -= 18

    pdf.save()

    return response


@login_required
@admin_required
def reporte_ventas_excel(request):

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename="reporte_ventas.xlsx"'

    wb = Workbook()
    ws = wb.active

    ws.title = "Ventas"

    headers = [
        "ID",
        "Título",
        "Cliente",
        "Vendedor",
        "Monto",
        "Estado",
        "Fecha creación",
        "Fecha cierre estimada",
    ]

    ws.append(headers)

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ventas = OportunidadVenta.objects.select_related(
        'cliente',
        'vendedor'
    ).all()

    for venta in ventas:

        ws.append([
            venta.id,
            venta.titulo,
            venta.cliente.nombre,
            venta.vendedor.username if venta.vendedor else "Sin asignar",
            float(venta.monto),
            venta.get_estado_display(),
            venta.fecha_creacion.strftime("%d/%m/%Y %H:%M"),
            venta.fecha_cierre_estimada.strftime("%d/%m/%Y")
            if venta.fecha_cierre_estimada else "No definida",
        ])

    for column_cells in ws.columns:

        length = max(
            len(str(cell.value)) if cell.value else 0
            for cell in column_cells
        )

        ws.column_dimensions[
            column_cells[0].column_letter
        ].width = length + 4

    wb.save(response)

    return response


@login_required
def reporte_seguimientos_pdf(request):

    if request.user.rol not in ['VENDEDOR', 'ADMIN']:

        messages.error(
            request,
            'No tienes permisos para descargar reportes.'
        )

        return redirect('seguimiento_list')

    response = HttpResponse(content_type='application/pdf')

    response['Content-Disposition'] = (
        'attachment; filename="reporte_seguimientos.pdf"'
    )

    pdf = canvas.Canvas(response, pagesize=letter)

    width, height = letter

    pdf.setTitle("Reporte de Seguimientos")

    pdf.setFillColor(colors.HexColor("#2563eb"))

    pdf.rect(
        0,
        height - 90,
        width,
        90,
        fill=True,
        stroke=False
    )

    pdf.setFillColor(colors.white)

    pdf.setFont("Helvetica-Bold", 22)

    pdf.drawString(
        40,
        height - 45,
        "Reporte de Seguimientos"
    )

    pdf.setFont("Helvetica", 11)

    pdf.drawString(
        40,
        height - 65,
        "Reporte generado correctamente desde GestionVentas CRM"
    )

    y = height - 120

    pdf.setFillColor(colors.HexColor("#0f172a"))

    pdf.setFont("Helvetica-Bold", 10)

    pdf.drawString(40, y, "ID")
    pdf.drawString(75, y, "Cliente")
    pdf.drawString(190, y, "Venta")
    pdf.drawString(310, y, "Tipo")
    pdf.drawString(395, y, "Usuario")
    pdf.drawString(480, y, "Estado")

    y -= 18

    pdf.setFont("Helvetica", 8)

    seguimientos = Seguimiento.objects.select_related(
        'cliente',
        'oportunidad',
        'usuario'
    ).all().order_by('-fecha_contacto')

    if request.user.rol == 'VENDEDOR':
        seguimientos = seguimientos.filter(
            usuario=request.user
        )

    for seguimiento in seguimientos:

        if y < 60:

            pdf.showPage()

            y = height - 50

            pdf.setFont("Helvetica-Bold", 10)

            pdf.drawString(40, y, "ID")
            pdf.drawString(75, y, "Cliente")
            pdf.drawString(190, y, "Venta")
            pdf.drawString(310, y, "Tipo")
            pdf.drawString(395, y, "Usuario")
            pdf.drawString(480, y, "Estado")

            y -= 18

            pdf.setFont("Helvetica", 8)

        estado = (
            "Completado"
            if seguimiento.completado
            else "Pendiente"
        )

        pdf.drawString(40, y, str(seguimiento.id))

        pdf.drawString(
            75,
            y,
            seguimiento.cliente.nombre[:22]
        )

        pdf.drawString(
            190,
            y,
            seguimiento.oportunidad.titulo[:22]
        )

        pdf.drawString(
            310,
            y,
            seguimiento.get_tipo_contacto_display()[:14]
        )

        pdf.drawString(
            395,
            y,
            seguimiento.usuario.username[:14]
        )

        pdf.drawString(
            480,
            y,
            estado
        )

        y -= 16

    pdf.setFillColor(colors.HexColor("#16a34a"))

    pdf.setFont("Helvetica-Bold", 12)

    pdf.drawString(
        40,
        40,
        "Reporte realizado correctamente."
    )

    pdf.save()

    return response